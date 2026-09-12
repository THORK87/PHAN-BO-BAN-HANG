import base64
import requests
import streamlit as st
import pandas as pd
import numpy as np
import io
import math
import secrets
from datetime import datetime, date
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Phân Bổ Bán Hàng Chuẩn Excel", layout="wide")

# =========================================================================
# CẤU HÌNH GITHUB DATABASE BẢN QUYỀN
# =========================================================================
GITHUB_REPO = st.secrets.get("GITHUB_REPO", "THORK87/PHAN-BO-BAN-HANG")
GITHUB_BRANCH = st.secrets.get("GITHUB_BRANCH", "main")
FILE_PATH = "licenses.json"
ADMIN_KEY = "Duykhuong@2026"  # Mật khẩu quản trị của bạn

# Lấy token từ Secrets của Streamlit Cloud
GITHUB_TOKEN = st.secrets.get("GITHUB_TOKEN", "")

HEADERS = {
    "Authorization": f"token {GITHUB_TOKEN}",
    "Accept": "application/vnd.github.v3+json"
}
API_URL = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{FILE_PATH}"

def get_remote_licenses():
    try:
        res = requests.get(f"{API_URL}?ref={GITHUB_BRANCH}", headers=HEADERS)
        if res.status_code == 200:
            data = res.json()
            content = base64.b64decode(data["content"]).decode("utf-8")
            return json.loads(content), data["sha"]
    except Exception:
        pass
    return {}, None

import json

def update_remote_licenses(new_data, sha=None):
    try:
        content_str = json.dumps(new_data, ensure_ascii=False, indent=4)
        content_b64 = base64.b64encode(content_str.encode("utf-8")).decode("utf-8")
        payload = {
            "message": f"Update licenses.json via Admin - {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            "content": content_b64,
            "branch": GITHUB_BRANCH
        }
        if sha:
            payload["sha"] = sha
        res = requests.put(API_URL, headers=HEADERS, json=payload)
        return res.status_code in [200, 201]
    except Exception:
        return False

# Đọc dữ liệu từ GitHub
db_licenses, file_sha = get_remote_licenses()

# =========================================================================
# 1. KIỂM TRA BẢN QUYỀN KHÁCH HÀNG (GIỮ NGUYÊN KEY)
# =========================================================================
if "is_licensed" not in st.session_state:
    st.session_state.is_licensed = False
if "is_admin" not in st.session_state:
    st.session_state.is_admin = False

if not st.session_state.is_licensed and not st.session_state.is_admin:
    st.sidebar.title("🔐 KÍCH HOẠT BẢN QUYỀN")
    user_key = st.sidebar.text_input("Nhập mã License Key:", type="password")
    
    if st.sidebar.button("Kích hoạt"):
        key_input = user_key.strip()
        if key_input == ADMIN_KEY:
            st.session_state.is_admin = True
            st.rerun()
            
        client_key = key_input.upper()
        if client_key in db_licenses:
            client_info = db_licenses[client_key]
            if client_info.get("status") == "blocked":
                st.sidebar.error("❌ Mã bản quyền này đã bị thu hồi hoặc khóa truy cập!")
            else:
                exp_date = datetime.strptime(client_info["expiry"], "%Y-%m-%d").date()
                if datetime.now().date() > exp_date:
                    st.sidebar.error(f"❌ Bản quyền đã hết hạn vào ngày {exp_date.strftime('%d/%m/%Y')}!")
                else:
                    st.session_state.is_licensed = True
                    st.sidebar.success(f"Hợp lệ! Hạn sử dụng: {exp_date.strftime('%d/%m/%Y')}")
                    st.rerun()
        else:
            st.sidebar.error("Mã kích hoạt không tồn tại trên hệ thống!")

    st.warning("⚠️ Vui lòng nhập License Key ở thanh menu bên trái để mở khóa phần mềm.")
    st.stop()

# =========================================================================
# 2. BẢNG ĐIỀU KHIỂN QUẢN TRỊ ADMIN (TẠO - GIA HẠN - HỦY)
# =========================================================================
if st.session_state.is_admin:
    st.title("🔑 BẢNG QUẢN TRỊ BẢN QUYỀN (ĐỒNG BỘ GITHUB)")
    st.info("Dữ liệu được lưu vĩnh viễn vào file licenses.json trên GitHub Repository.")
    
    tab_tao, tab_ql = st.tabs(["➕ Tạo Key Mới (Không ngày tháng)", "📋 Danh Sách & Gia Hạn / Hủy"])

   # TAB 1: TẠO KEY MỚI CHUẨN QUỐC TẾ (4 CỤM ĐỘC LẬP)
    with tab_tao:
        c1, c2 = st.columns(2)
        with c1:
            ten_khach = st.text_input("Ghi chú tên khách hàng (quản lý nội bộ):", value="CONGTY_ABC")
        with c2:
            ngay_het = st.date_input("Hạn sử dụng ban đầu:", value=date(2027, 1, 1))
            
        if st.button("🚀 Tạo License Key Quốc Tế", type="primary"):
            c_name = ten_khach.strip() if ten_khach.strip() else "KHACH_HANG"
            
            # Tạo 4 cụm ngẫu nhiên: XXXX-XXXX-XXXX-XXXX
            parts = [secrets.token_hex(2).upper() for _ in range(4)]
            generated_key = "-".join(parts)
            
            db_licenses[generated_key] = {
                "client_name": c_name,
                "expiry": ngay_het.strftime("%Y-%m-%d"),
                "status": "active"
            }
            if update_remote_licenses(db_licenses, file_sha):
                st.success(f"Đã tạo Key thành công cho {c_name}!")
                st.code(generated_key, language="text")
                st.info("Gửi mã trên cho khách hàng. Khách sẽ dùng cố định mã này vĩnh viễn.")
                st.rerun()
            else:
                st.error("Lỗi khi lưu lên GitHub. Vui lòng kiểm tra lại GITHUB_TOKEN trong Secrets.")

    # TAB 2: QUẢN LÝ - GIA HẠN - HỦY - XÓA
    with tab_ql:
        if not db_licenses:
            st.info("Chưa có mã bản quyền nào trên GitHub.")
        else:
            for k, v in list(db_licenses.items()):
                ten_hien_thi = v.get("client_name", k)
                with st.expander(f"Khách hàng: {ten_hien_thi} | Key: {k}", expanded=True):
                    col_info, col_han, col_action = st.columns([2.5, 2, 2])
                    
                    with col_info:
                        st.write(f"**Khách hàng:** `{ten_hien_thi}`")
                        # Hộp code có sẵn nút Copy nhỏ ở góc phải
                        st.code(k, language="text")
                        st.write(f"Trạng thái: {'🟢 Hoạt động' if v['status'] == 'active' else '🔴 ĐÃ KHÓA'}")
                        st.write(f"Hạn: **{datetime.strptime(v['expiry'], '%Y-%m-%d').strftime('%d/%m/%Y')}**")
                        
                    with col_han:
                        cur_date = datetime.strptime(v['expiry'], "%Y-%m-%d").date()
                        new_date = st.date_input("Chọn hạn mới:", value=cur_date, key=f"date_{k}")
                        if st.button("Cập nhật hạn", key=f"btn_date_{k}"):
                            db_licenses[k]["expiry"] = new_date.strftime("%Y-%m-%d")
                            if update_remote_licenses(db_licenses, file_sha):
                                st.success("Đã gia hạn thành công!")
                                st.rerun()
                            else:
                                st.error("Lỗi khi cập nhật lên GitHub!")
                                
                    with col_action:
                        st.write("Thao tác bản quyền:")
                        # Nút Khóa / Mở khóa
                        if v['status'] == "active":
                            if st.button("🚫 Khóa Key", key=f"block_{k}", type="secondary"):
                                db_licenses[k]["status"] = "blocked"
                                update_remote_licenses(db_licenses, file_sha)
                                st.warning("Đã khóa bản quyền!")
                                st.rerun()
                        else:
                            if st.button("✅ Mở khóa", key=f"unblock_{k}"):
                                db_licenses[k]["status"] = "active"
                                update_remote_licenses(db_licenses, file_sha)
                                st.success("Đã mở khóa lại!")
                                st.rerun()
                        
                        st.write("")
                        # Nút Xóa key hoàn toàn khỏi GitHub Database
                        if st.button("🗑️ Xóa vĩnh viễn", key=f"del_{k}", type="primary"):
                            del db_licenses[k]
                            if update_remote_licenses(db_licenses, file_sha):
                                st.success(f"Đã xóa vĩnh viễn key `{k}`!")
                                st.rerun()
                            else:
                                st.error("Lỗi khi xóa key trên GitHub!")
# =========================================================================
# 1. NẠP DỮ LIỆU TỪ SHEET T4.2026 (ĐƠN GIÁ LÀM TRÒN SỐ NGUYÊN)
# =========================================================================
st.subheader("1. DANH MỤC HÀNG HÓA")

# Tạo file Excel mẫu chuẩn
df_mau = pd.DataFrame({
    "Mã VTHH (*)": ["SP01", "SP02", "SP03"],
    "Tên VTHH (*)": ["Sản phẩm mẫu A", "Sản phẩm mẫu B", "Sản phẩm mẫu C"],
    "ĐVT chính": ["Cái", "Hộp", "Gói"],
    "Giá bán cố định": [50000, 120000, 25000],
    "Tồn gốc (ẩn)": [100, 50, 200]
})
buffer_mau = io.BytesIO()
with pd.ExcelWriter(buffer_mau, engine="openpyxl") as writer:
    df_mau.to_excel(writer, sheet_name="T4.2026", index=False)

col_up, col_btn = st.columns([3, 1])
with col_up:
    file_upload = st.file_uploader("Kéo thả file số liệu vào đây:", type=["xlsx"])
with col_btn:
    st.write("")
    st.write("")
    st.download_button(
        label="📥 Tải file mẫu Excel",
        data=buffer_mau.getvalue(),
        file_name="MAU_PHAN_BO_BAN_HANG.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        help="Tải file mẫu Excel có sẵn các cột chuẩn để nhập liệu"
    )

if file_upload is not None:
    try:
        xls = pd.ExcelFile(file_upload)
        sheet_target = "T4.2026" if "T4.2026" in xls.sheet_names else xls.sheet_names[0]
        df_source = pd.read_excel(xls, sheet_target)
        df_source = df_source.dropna(subset=[df_source.columns[0]]).copy()

        df_init = pd.DataFrame({
            "MÃ VTHH": df_source["Mã VTHH (*)"],
            "TÊN VTHH": df_source["Tên VTHH (*)"],
            "ĐVT": df_source["ĐVT chính"],
            "ĐƠN GIÁ": pd.to_numeric(df_source["Giá bán cố định"], errors="coerce").fillna(0).round(0).astype(int),
            "TỒN ĐẦU": pd.to_numeric(df_source["Tồn gốc (ẩn)"], errors="coerce").fillna(0).round(0).astype(int)
        })
        st.success(f"Đã nạp chính xác {len(df_init)} mặt hàng từ sheet '{sheet_target}'.")
    except Exception as e:
        st.error(f"Lỗi đọc file: {e}")
        df_init = pd.DataFrame(columns=["MÃ VTHH", "TÊN VTHH", "ĐVT", "ĐƠN GIÁ", "TỒN ĐẦU"])
else:
    df_init = pd.DataFrame([
        {"MÃ VTHH": "Bánh Khoai Tây xá 4kg", "TÊN VTHH": "Bánh Khoai Tây xá 4kg", "ĐVT": "kg", "ĐƠN GIÁ": 58850, "TỒN ĐẦU": 4},
        {"MÃ VTHH": "VP26_01", "TÊN VTHH": "Nước T-REXX Oishi Túi 180ml (5 bịch/T)", "ĐVT": "Bịch", "ĐƠN GIÁ": 28160, "TỒN ĐẦU": 25},
        {"MÃ VTHH": "VP26_02", "TÊN VTHH": "Bánh snack Oishi 2K (10 bịch/bao)", "ĐVT": "Bịch", "ĐƠN GIÁ": 15248, "TỒN ĐẦU": 280},
        {"MÃ VTHH": "VP26_10", "TÊN VTHH": "Bánh snack Oishi 5K (8 bịch/bao)", "ĐVT": "Bịch", "ĐƠN GIÁ": 38500, "TỒN ĐẦU": 8},
        {"MÃ VTHH": "VP26_13", "TÊN VTHH": "Bánh Chấm 330g PC (12 cái)", "ĐVT": "Cái", "ĐƠN GIÁ": 2383, "TỒN ĐẦU": 120},
        {"MÃ VTHH": "VP26_32", "TÊN VTHH": "Thạch miếng nhỏ Suso LƯỚI 800g (62 cái)", "ĐVT": "Cái", "ĐƠN GIÁ": 550, "TỒN ĐẦU": 372},
        {"MÃ VTHH": "VP26_40", "TÊN VTHH": "Coke 300ml - Coca Chai Nhí 5K", "ĐVT": "Chai", "ĐƠN GIÁ": 4285, "TỒN ĐẦU": 24},
        {"MÃ VTHH": "VP26_41", "TÊN VTHH": "Đậu phộng Oishi 1K (20 dây/T)", "ĐVT": "Dây", "ĐƠN GIÁ": 7637, "TỒN ĐẦU": 20},
        {"MÃ VTHH": "VP26_43", "TÊN VTHH": "Thạch SUSO Cây Dài (24 cái)", "ĐVT": "Cái", "ĐƠN GIÁ": 1375, "TỒN ĐẦU": 432},
        {"MÃ VTHH": "VP26_45", "TÊN VTHH": "Mì Vàng Enaak 30g (24 gói/Hộp)", "ĐVT": "Gói", "ĐƠN GIÁ": 5500, "TỒN ĐẦU": 144},
        {"MÃ VTHH": "VP26_46", "TÊN VTHH": "Nước ÉP THẠCH Suso (48 ly)", "ĐVT": "Ly", "ĐƠN GIÁ": 3942, "TỒN ĐẦU": 480},
        {"MÃ VTHH": "VP26_47", "TÊN VTHH": "Kẹo Cao Su Dinos 500g (69 viên)", "ĐVT": "Viên", "ĐƠN GIÁ": 829, "TỒN ĐẦU": 759},
        {"MÃ VTHH": "VP26_49", "TÊN VTHH": "Bánh Mứt Dâu Xá 4kg PC", "ĐVT": "kg", "ĐƠN GIÁ": 61875, "TỒN ĐẦU": 4},
        {"MÃ VTHH": "VP26_50", "TÊN VTHH": "Nước Túi Oishi (10 túi x 5 bịch / T)", "ĐVT": "Túi", "ĐƠN GIÁ": 2560, "TỒN ĐẦU": 500},
    ])

edited_df = st.data_editor(
    df_init,
    num_rows="dynamic",
    use_container_width=True,
    column_config={
        "ĐƠN GIÁ": st.column_config.NumberColumn("ĐƠN GIÁ (VNĐ)", format="%d"),
        "TỒN ĐẦU": st.column_config.NumberColumn("TỒN ĐẦU", format="%d"),
    }
)

# =========================================================================
# 2. THIẾT LẬP THÔNG SỐ PHÂN BỔ
# =========================================================================
st.subheader("2. THIẾT LẬP THÔNG SỐ PHÂN BỔ")
c1, c2, c3 = st.columns(3)

with c1:
    so_ngay = st.number_input("SỐ NGÀY BÁN (Nhập số nguyên):", min_value=1, max_value=31, value=5, step=1)
with c2:
    bien_do = st.number_input("ĐỘ LỆCH DOANH THU CÁC NGÀY (0.0 = san phẳng tuyệt đối):", min_value=0.0, max_value=3.0, value=0.0, step=0.05, format="%.2f")
with c3:
    tong_tien_muc_tieu = st.number_input("DOANH THU TỔNG (VNĐ):", min_value=0.0, value=10000000.0, step=500000.0, format="%.0f")
    st.caption(f"👉 Số tiền: **{int(tong_tien_muc_tieu):,} đ**")

btn_run = st.button("🚀 CHẠY PHÂN BỔ ", type="primary")

# =========================================================================
# 3. THUẬT TOÁN TỐI ƯU 2 CHIỀU (KNAPSACK GREEDY)
# =========================================================================
if btn_run:
    df_items = edited_df.dropna(subset=["MÃ VTHH"]).copy().reset_index(drop=True)
    df_items["ĐƠN GIÁ"] = pd.to_numeric(df_items["ĐƠN GIÁ"], errors="coerce").fillna(0).round(0).astype(int)
    df_items["TỒN ĐẦU"] = pd.to_numeric(df_items["TỒN ĐẦU"], errors="coerce").fillna(0).round(0).astype(int)
    
    n = len(df_items)
    if n == 0 or tong_tien_muc_tieu <= 0:
        st.warning("Vui lòng kiểm tra lại danh sách hàng và số tiền phân bổ!")
        st.stop()

    so_ngay_int = int(so_ngay)
    danh_sach_ngay = [f"Ngày {d}" for d in range(1, so_ngay_int + 1)]

    ton_arr = df_items["TỒN ĐẦU"].values
    gia_arr = df_items["ĐƠN GIÁ"].values

    # --- TÍNH BU: LẤY ĐA DẠNG HÀNG HÓA KHI SỐ TIỀN NHỎ ---
    tong_gia_tri_ton = np.sum(ton_arr * gia_arr)
    bu_arr = np.zeros(n, dtype=int)
    tien_hien_tai = 0.0

    idx_sorted_by_price = np.argsort(gia_arr)
    for i in idx_sorted_by_price:
        if gia_arr[i] > 0 and ton_arr[i] > 0:
            if tien_hien_tai + gia_arr[i] <= tong_tien_muc_tieu:
                bu_arr[i] = 1
                tien_hien_tai += gia_arr[i]

    tien_con_lai = tong_tien_muc_tieu - tien_hien_tai
    ton_con_lai = ton_arr - bu_arr
    tong_ton_con_lai_vnd = np.sum(ton_con_lai * gia_arr)

    if tien_con_lai > 0 and tong_ton_con_lai_vnd > 0:
        ti_le_phu = min(1.0, tien_con_lai / tong_ton_con_lai_vnd)
        sl_them = np.floor(ton_con_lai * ti_le_phu).astype(int)
        sl_them = np.minimum(sl_them, ton_con_lai)
        bu_arr += sl_them
        tien_hien_tai = np.sum(bu_arr * gia_arr)

    tien_con_lai = tong_tien_muc_tieu - tien_hien_tai
    if tien_con_lai > 0:
        for i in range(n):
            if gia_arr[i] > 0 and bu_arr[i] < ton_arr[i]:
                sl_bu_max = ton_arr[i] - bu_arr[i]
                sl_can_bu = int(tien_con_lai // gia_arr[i])
                sl_thuc_bu = min(sl_bu_max, sl_can_bu)
                if sl_thuc_bu > 0:
                    bu_arr[i] += sl_thuc_bu
                    tien_con_lai -= sl_thuc_bu * gia_arr[i]
                    if tien_con_lai < np.min(gia_arr[gia_arr > 0]):
                        break

    df_items["TỔNG SL BÁN"] = bu_arr
    df_items["THÀNH TIỀN"] = bu_arr * gia_arr
    df_items["TỒN CUỐI"] = ton_arr - bu_arr

   # 1. Xác định doanh thu mục tiêu từng ngày (theo bien_do)
    avg_day = tong_tien_muc_tieu / so_ngay_int
    day_targets = np.zeros(so_ngay_int, dtype=float)
    for d in range(so_ngay_int):
        wave = np.sin((d + 1) * 2 * np.pi / so_ngay_int) if so_ngay_int > 1 else 0
        day_targets[d] = avg_day * (1.0 + bien_do * wave)
    
    day_targets = day_targets * (tong_tien_muc_tieu / np.sum(day_targets))
    day_ratios = day_targets / np.sum(day_targets)  # Tỷ trọng mục tiêu từng ngày

    # 2. Phân bổ ma trận đồng đều đa món (Mọi món cùng co dãn theo biên độ)
    matrix_ngay = np.zeros((n, so_ngay_int), dtype=int)

    for i in range(n):
        total_qty = bu_arr[i]
        if total_qty == 0:
            continue
            
        # Món số lượng ít (<= 3): Rải đều cách nhật để không mất món ở các ngày
        if total_qty <= 3:
            indices = [int(round(k * (so_ngay_int - 1) / (total_qty - 1))) if total_qty > 1 else so_ngay_int // 2 for k in range(total_qty)]
            for d_idx in indices:
                matrix_ngay[i, d_idx] += 1
        else:
            # Món số lượng lớn: Co dãn theo tỷ trọng ngày để mọi món cùng gánh biên độ
            raw_dist = total_qty * day_ratios
            base_dist = np.floor(raw_dist).astype(int)
            matrix_ngay[i, :] += base_dist
            
            # Phân phối phần dư lẻ vào ngày có phần thập phân cao nhất
            rem = total_qty - np.sum(base_dist)
            if rem > 0:
                fractional_parts = raw_dist - base_dist
                top_days = np.argsort(-fractional_parts)[:rem]
                for d_idx in top_days:
                    matrix_ngay[i, d_idx] += 1
    # Bước C: Tạo DataFrame kết quả hiển thị
    df_daily = pd.DataFrame(matrix_ngay, columns=danh_sach_ngay)

    df_ket_qua_chi_tiet = pd.concat([
        df_items[["MÃ VTHH", "TÊN VTHH", "ĐVT", "ĐƠN GIÁ", "TỒN ĐẦU"]],
        df_daily,
        df_items[["TỔNG SL BÁN", "THÀNH TIỀN", "TỒN CUỐI"]]
    ], axis=1)

    # =========================================================================
    # 4. HIỂN THỊ KẾT QUẢ & XUẤT FILE EXCEL
    # =========================================================================
    st.divider()
    st.subheader(f"3. KẾT QUẢ PHÂN BỔ (TỔNG CỘNG {so_ngay_int} NGÀY)")

    tong_tien_thuc_te = int(np.sum(df_items["THÀNH TIỀN"]))
    chenh_lech = tong_tien_thuc_te - int(tong_tien_muc_tieu)

    m1, m2, m3 = st.columns(3)
    m1.metric("DOANH THU MỤC TIÊU", f"{tong_tien_muc_tieu:,.0f} đ")
    m2.metric("ĐÃ PHÂN BỔ THỰC TẾ", f"{tong_tien_thuc_te:,.0f} đ")
    m3.metric("CHÊNH LỆCH", f"{chenh_lech:,.0f} đ")

    st.markdown("#### 💰 TỔNG CỘNG DOANH THU VÀ CHỈ TIÊU TỪNG NGÀY")
    
    doanh_thu_tung_ngay = np.dot(gia_arr, matrix_ngay)
    tong_sl_tung_ngay = np.sum(matrix_ngay, axis=0)
    so_mon_tung_ngay = np.sum(matrix_ngay > 0, axis=0)

    df_tong_hop_ngay = pd.DataFrame({
        "CHỈ TIÊU": [
            "DOANH THU TỪNG NGÀY (VNĐ)",
            "Tổng số lượng bán (ĐVT)",
            "Số chủng loại mặt hàng bán"
        ]
    })
    for idx, d_col in enumerate(danh_sach_ngay):
        df_tong_hop_ngay[d_col] = [
            f"{doanh_thu_tung_ngay[idx]:,.0f} đ",
            f"{tong_sl_tung_ngay[idx]:,d}",
            f"{so_mon_tung_ngay[idx]} món"
        ]
    df_tong_hop_ngay["TỔNG CỘNG"] = [
        f"{tong_tien_thuc_te:,.0f} đ",
        f"{np.sum(bu_arr):,d}",
        f"{np.sum(bu_arr > 0)} món"
    ]

    st.dataframe(df_tong_hop_ngay, use_container_width=True, hide_index=True)

    st.markdown("#### 📋 CHI TIẾT SẢN LƯỢNG BÁN THEO MẶT HÀNG")
    st.dataframe(
        df_ket_qua_chi_tiet,
        use_container_width=True,
        hide_index=True,
        column_config={
            "ĐƠN GIÁ": st.column_config.NumberColumn("ĐƠN GIÁ", format="%,d"),
            "THÀNH TIỀN": st.column_config.NumberColumn("THÀNH TIỀN", format="%,d")
        }
    )

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_ket_qua_chi_tiet.to_excel(writer, sheet_name="SL_BAN_THEO_NGAY", index=False)

        ws = writer.sheets["SL_BAN_THEO_NGAY"]
        
        last_row = ws.max_row
        row_sl_idx = last_row + 1
        row_tien_idx = last_row + 2

        ws.cell(row_sl_idx, 1, "TỔNG CỘNG")
        ws.cell(row_sl_idx, 2, "Tổng số lượng bán")
        ws.cell(row_sl_idx, 5, int(np.sum(ton_arr)))
        for idx in range(so_ngay_int):
            ws.cell(row_sl_idx, 6 + idx, int(tong_sl_tung_ngay[idx]))
        ws.cell(row_sl_idx, 6 + so_ngay_int, int(np.sum(bu_arr)))
        ws.cell(row_sl_idx, 7 + so_ngay_int, int(tong_tien_thuc_te))
        ws.cell(row_sl_idx, 8 + so_ngay_int, int(np.sum(ton_arr - bu_arr)))

        ws.cell(row_tien_idx, 1, "DOANH THU")
        ws.cell(row_tien_idx, 2, "Tổng tiền bán ngày")
        for idx in range(so_ngay_int):
            ws.cell(row_tien_idx, 6 + idx, int(doanh_thu_tung_ngay[idx]))
        ws.cell(row_tien_idx, 7 + so_ngay_int, int(tong_tien_thuc_te))

        header_font = Font(name="Arial", size=11, bold=True, color="000000")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

        sum_font = Font(name="Arial", size=11, bold=True, color="002060")
        sum_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = header_alignment
            cell.fill = header_fill
        ws.row_dimensions[1].height = 28

        for r in [row_sl_idx, row_tien_idx]:
            for cell in ws[r]:
                cell.font = sum_font
                cell.fill = sum_fill
            ws.row_dimensions[r].height = 24

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    st.download_button(
        label="📥 TẢI FILE EXCEL KẾT QUẢ (KHỚP CHUẨN FILE GỐC)",
        data=output.getvalue(),
        file_name=f"PHAN_BO_{so_ngay_int}_NGAY.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
